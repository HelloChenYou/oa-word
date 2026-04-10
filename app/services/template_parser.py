import json
import re
from pathlib import Path

from docx import Document


PLACEHOLDER_PATTERN = re.compile(r"\{([a-zA-Z0-9_]+)\}")


def parse_template_text(raw_text: str) -> dict:
    placeholders = list(dict.fromkeys(PLACEHOLDER_PATTERN.findall(raw_text)))
    lines = [line.strip() for line in raw_text.splitlines() if line.strip()]
    sections = []
    for line in lines:
        if line.startswith("【") and "】" in line:
            sections.append(line.split("】", 1)[0] + "】")
    return {
        "placeholders": placeholders,
        "required_sections": list(dict.fromkeys(sections)),
    }


def read_template_content(file_path: Path, file_type: str) -> str:
    if file_type in {"txt", "md"}:
        return file_path.read_text(encoding="utf-8")
    if file_type == "docx":
        doc = Document(str(file_path))
        return "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
    raise ValueError(f"unsupported file type: {file_type}")


def parse_template_file(file_path: Path, file_type: str) -> tuple[str, str]:
    raw_text = read_template_content(file_path, file_type)
    parsed = parse_template_text(raw_text)
    return raw_text, json.dumps(parsed, ensure_ascii=False)
